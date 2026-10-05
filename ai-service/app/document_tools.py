"""Text extraction tools for private company knowledge documents."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_document(path: Path, *, max_bytes: int = 10 * 1024 * 1024,
                     max_pages: int = 200, max_characters: int = 500_000) -> list[dict]:
    """Return page/section text; reject unsupported, oversized, or empty files."""
    path = Path(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported document type")
    if path.stat().st_size > max_bytes:
        raise ValueError("Document exceeds the configured size limit")

    if extension == ".pdf":
        reader = PdfReader(path)
        if len(reader.pages) > max_pages:
            raise ValueError("PDF exceeds the configured page limit")
        sections = [
            {"page_number": number, "section": None, "text": page.extract_text() or ""}
            for number, page in enumerate(reader.pages, start=1)
        ]
    elif extension == ".docx":
        document = Document(path)
        sections = [{"page_number": None, "section": "Document", "text": "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )}]
    else:
        sections = [{"page_number": None, "section": "Document", "text": path.read_text(encoding="utf-8-sig")}]

    result = []
    total = 0
    for section in sections:
        content = section["text"].strip()
        total += len(content)
        if total > max_characters:
            raise ValueError("Extracted text exceeds the configured character limit")
        if content:
            result.append({**section, "text": content})
    if not result:
        raise ValueError("No extractable text found; scanned documents need OCR")
    return result
